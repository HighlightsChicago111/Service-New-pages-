"""Merge a content batch into the V2 workbook (replaces rows for the same service IDs).

Usage: python -I scripts/content/add_batch.py "<Content Plan.xlsx>"
"""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.dirname(HERE))

from openpyxl import load_workbook  # noqa: E402

from batch_341_345 import IMG, ITEMS, PAGES  # noqa: E402
from workbook_spec import PAGE_COLUMNS  # noqa: E402

WORKBOOK = "Highlights Chicago - Service Pages V2 Data.xlsx"

# Content Plan '3 Service Level' review assignments for this batch.
REVIEW_IDS = {
    "341": ["R019", "R090", "R080", "R001"],
    "342": ["R073", "R083", "R048", "R096"],
    "343": ["R134", "R033", "R108", "R091"],
    "344": ["R068", "R034", "R161", "R038"],
    "345": ["R055", "R053", "R056", "R063"],
}

# Related local covers, labelled illustrative until real job photos exist.
IMAGES = {
    "341": [("solar-panel-installation.jpg", "solar panel installation"), ("solar-battery-installation.jpg", "solar battery installation"), ("electrical-panel-upgrade.jpg", "electrical panel upgrade")],
    "342": [("low-voltage-wiring-installation.jpg", "low voltage wiring installation"), ("structured-cabling-installation.jpg", "structured cabling installation"), ("electrical-outlet-installation.jpg", "electrical outlet installation")],
    "343": [("electrical-installation-services.jpg", "electrical installation work"), ("electrical-panel-upgrade.jpg", "electrical panel upgrade"), ("energy-star-appliances.jpg", "appliance circuit work")],
    "344": [("light-switch-replacement-and-installation.jpg", "light switch and dimmer installation"), ("recessed-lighting-installation.jpg", "recessed lighting installation"), ("light-fixture-installation-and-replacement.jpg", "light fixture installation")],
    "345": [("fire-alarm-installation.jpg", "life-safety device installation"), ("electrical-installation-services.jpg", "commercial electrical installation"), ("outdoor-lighting-installation.jpg", "exterior lighting installation")],
}


def load_reviews(plan_path):
    wanted = {rid for ids in REVIEW_IDS.values() for rid in ids}
    wb = load_workbook(plan_path, read_only=True, data_only=True)
    rows = wb["Reviews"].iter_rows(values_only=True)
    header = None
    found = {}
    for row in rows:
        if header is None:
            if row and row[0] == "Review ID":
                header = list(row)
            continue
        if row and row[0] in wanted:
            record = dict(zip(header, row))
            words = str(record["Review"]).split()
            excerpt = " ".join(words[:14]) + ("…" if len(words) > 14 else "")
            found[row[0]] = [row[0], record["Reviewer"], record["ISO Date"], record["Review Link"], excerpt, record.get("Named Location") or ""]
    missing = wanted - set(found)
    if missing:
        raise SystemExit(f"Reviews not found in Content Plan: {sorted(missing)}")
    return found


def main():
    plan = sys.argv[1]
    reviews = load_reviews(plan)
    ids = {p["service_id"] for p in PAGES}
    wb = load_workbook(WORKBOOK)
    pages_ws, items_ws = wb["01 Pages"], wb["02 Items"]

    cols = [c.value for c in pages_ws[1]]
    for row in range(pages_ws.max_row, 1, -1):
        if str(pages_ws.cell(row, 1).value) in ids:
            pages_ws.delete_rows(row)
    for row in range(items_ws.max_row, 1, -1):
        if str(items_ws.cell(row, 1).value) in ids:
            items_ws.delete_rows(row)

    expected = [c[0] for c in PAGE_COLUMNS]
    for page in PAGES:
        unknown = set(page) - set(expected)
        if unknown:
            raise SystemExit(f"{page['service_id']}: unknown columns {unknown}")
        pages_ws.append([page.get(c, "") for c in cols])

    for page in PAGES:
        sid = page["service_id"]
        rows = list(ITEMS.get(sid, []))
        rows += [("review", reviews[rid]) for rid in REVIEW_IDS[sid]]
        for block in ("gallery", "working_photo"):
            for file, subject in IMAGES[sid]:
                alt = f"Illustrative photo of {subject} by Highlights Chicago electricians in Chicago"
                rows.append((block, [IMG + file, alt, f"Illustrative photo of {subject}"]))
        counters = {}
        for block, values in rows:
            counters[block] = counters.get(block, 0) + 1
            items_ws.append([sid, block, counters[block], *(list(values) + [""] * 6)[:6]])

    wb.save(WORKBOOK)
    print(f"Merged {len(PAGES)} pages ({', '.join(sorted(ids))}) into {WORKBOOK}")


if __name__ == "__main__":
    main()
