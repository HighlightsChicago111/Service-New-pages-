"""Build the V2 page-data workbook.

Creates "Highlights Chicago - Service Pages V2 Data.xlsx" with every sheet the
V2 writing skill fills, a Field Map that ties each column to its Sanity field,
and one complete example page (302, Solar) taken from the approved design file.

Usage (run from the project root):
  python -I scripts/make_workbook.py --template "<Service Landing Page new Template.html>" \
      --source "<old project>/data/source-content.json" [--out "<file>.xlsx"]
"""

import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from bs4 import BeautifulSoup, NavigableString  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from workbook_spec import (  # noqa: E402
    ANCHORS, AREA_BLOCKS, AREA_COLUMNS, ITEM_COLUMNS, PAGE_BLOCKS, PAGE_COLUMNS, SETTINGS_BLOCKS, SETTINGS_FIELDS,
)

DEFAULT_OUT = "Highlights Chicago - Service Pages V2 Data.xlsx"
HEADER_FILL = PatternFill("solid", fgColor="151F2A")
REQ_FILL = PatternFill("solid", fgColor="9EC837")
EXAMPLE_FILL = PatternFill("solid", fgColor="F4F6F7")


def text(node):
    return re.sub(r"\s+", " ", node.get_text(" ", strip=True)).replace(" ,", ",").replace(" .", ".").strip() if node else ""


def own_text(node):
    """Text of a node without its trailing link (link label/anchor are separate columns)."""
    if not node:
        return "", "", ""
    link = node.find("a")
    label = anchor = ""
    if link:
        label, anchor = text(link), anchor_of(link.get("href"))
        link.extract()
    return text(node), label, anchor


def ph(node):
    """Keep [CONFIRM: ...] placeholders as plain text."""
    for span in node.select(".nx-ph"):
        span.replace_with(NavigableString(span.get_text()))
    return node


# The design file used solar-specific ids; V2 uses generic ones for every service.
ANCHOR_RENAMES = {"panel-routes": "assessment-routes", "battery-scope": "equipment-options", "solar-repair": "diagnose", "your-solar-job": "your-job"}


def anchor_of(href):
    value = (href or "").lstrip("#")
    return ANCHOR_RENAMES.get(value, value)


def strip_area(value, area="Chicago"):
    return re.sub(rf"\s+in {area}\??$", "", value).strip()


def parse_template(path, source):
    soup = BeautifulSoup(open(path, encoding="utf-8").read(), "html.parser")
    for tag in soup.select("script, svg"):
        tag.decompose()
    main = soup.select_one("main.service-landing")
    ph(main)
    page, items = {}, []
    sid = "302"
    solar = next(row for row in source["equip"] if row["slug"] == "solar-panel-installation")
    solar_page = next(row for row in source["page"] if row["equipment_slug"] == "solar-panel-installation")

    def add(block, *values):
        order = sum(1 for item in items if item[1] == block) + 1
        items.append([sid, block, order, *values])

    crumb = text(main.select_one(".crumbs li[aria-current]"))
    page.update(
        service_id=sid, slug="solar-panel-installation", name=strip_area(crumb), parent_name=solar["parent_name"], area_slug="chicago",
        meta_title=strip_area(text(soup.title).replace(" | Highlights Chicago", "")) + " in Chicago",
        meta_description=soup.find("meta", attrs={"name": "description"})["content"],
        kw_primary=" || ".join(k.strip() for k in solar["kw_primary"].split(",")), kw_volume=solar["kw_volume"],
        kw_secondary=" || ".join(k.strip() for k in solar["kw_secondary"].split(",")),
    )
    hero = main.select_one("header.hero")
    page.update(h1_prefix=strip_area(text(hero.h1)), hero_lede=text(hero.select_one("p.lede")), cta_secondary=text(hero.select_one(".btn-secondary")))
    issue = hero.select_one("select[name=issue]")
    page.update(issue_question=text(hero.find("label", attrs={"for": issue["id"]})), issue_options=" || ".join(text(o) for o in issue.find_all("option")))
    for name in ("existingSystem", "symptom"):
        select = hero.select_one(f"select[name={name}]")
        add("form_field", name, text(hero.find("label", attrs={"for": select["id"]})), " || ".join(text(o) for o in select.find_all("option")), "", "", "")
    for img in hero.select(".cs-gallery img"):
        add("gallery", img["src"], img.get("alt", ""), img.get("title", ""), "", "", "")

    jobs = main.select_one("#your-solar-job")
    page.update(jobs_heading=text(jobs.h2), jobs_lede=text(jobs.select_one("p.lede")), jobs_note=text(jobs.select_one("p.section-note")))
    for card in jobs.select(".nx-path"):
        dds = card.find_all("dd")
        link = card.find("a")
        add("job_path", text(card.h3), text(card.p), text(dds[0]), text(dds[1]), text(link), anchor_of(link["href"]))

    equip = main.select_one("#equipment")
    note, label, anchor = own_text(equip.select_one("p.section-note"))
    page.update(equip_heading=text(equip.h2), equip_lede=text(equip.select_one("p.lede")), equip_footnote=note,
                equip_footnote_link_label=label, equip_footnote_link_anchor=anchor)
    seen = set()
    for tile in equip.select(".equip"):
        name = text(tile.b)
        if name not in seen:
            seen.add(name)
            add("equipment", name, text(tile.span), "", "", "", "")
    compare = equip.select_one("#equipment-choices")
    headers = [text(th) for th in compare.select("thead th")][1:]
    page.update(compare_heading=text(compare.h3), compare_intro=text(compare.select_one(".nx-sub")), compare_columns=" || ".join(headers))
    for row in compare.select("tbody tr"):
        cells = [text(td) for td in row.find_all("td")]
        add("compare_row", *(cells + [""] * 6)[:6])
    battery = equip.select_one("#battery-scope")
    callout = battery.select_one(".nx-callout")
    lead = text(callout.strong)
    callout.strong.extract()
    page.update(options_heading=text(battery.h3), options_callout_lead=lead, options_callout_body=text(callout), options_note=text(battery.select_one("p.small")))
    for option in battery.select(".nx-option"):
        add("option", text(option.h4), text(option.p), "", "", "", "")

    assess = main.select_one("#site-assessment")
    routes = assess.select_one("#panel-routes")
    page.update(assess_heading=text(assess.h2), assess_lede=text(assess.select_one("p.lede")), routes_heading=text(routes.h3), routes_note=text(routes.select_one("p.small")))
    for li in assess.select(".nx-checks li"):
        add("assess_check", text(li.b), text(li.span), "", "", "", "")
    for li in routes.select(".nx-routes li"):
        add("route", text(li.b), text(li.span), "", "", "", "")

    repair = main.select_one("#solar-repair")
    boxes = repair.select(".nx-two .nx-block")
    page.update(diagnose_heading=text(repair.h2), diagnose_lede=text(repair.select_one("p.lede")),
                diagnose_columns=" || ".join(text(th) for th in repair.select("thead th")),
                repair_title=text(boxes[0].h3), repair_body=text(boxes[0].p), orphan_title=text(boxes[1].h3), orphan_body=text(boxes[1].p))
    for row in repair.select("tbody tr"):
        cells = [text(td) for td in row.find_all("td")]
        add("symptom", *(cells + [""] * 6)[:6])

    brands = main.select_one("#brands")
    names = []
    for tile in brands.select(".brand-tile strong"):
        if text(tile) not in names:
            names.append(text(tile))
    page.update(brands_heading=strip_area(text(brands.h2)), brands_lede=text(brands.select_one("p.lede")), brands=" || ".join(names), brands_note=text(brands.select_one("p.section-note")))

    for card in main.select("#reviews .rev-card"):
        author = text(card.select_one(".rev-meta strong"))
        match = re.search(rf"::{re.escape(author)}::(\d{{4}}-\d{{2}}-\d{{2}})::[^:]+(?::[^:]+)*?::(R\d+)", solar_page["reviews"])
        words = text(card.blockquote).split()
        excerpt = " ".join(words[:14]) + ("…" if len(words) > 14 else "")
        add("review", match.group(2) if match else "", author, match.group(1) if match else "", card["href"], excerpt, "")

    why = main.select_one("#why-us")
    page.update(why_heading=text(why.h2), why_lede=text(why.select_one("p.lede")))
    for item in why.select(".why-item"):
        body, label, anchor = own_text(item.select_one(".why-body"))
        add("why", text(item.h3), body, label, anchor, "", "")

    for fig in main.select("#working-in-area figure"):
        img = fig.img
        add("working_photo", img["src"], img.get("alt", ""), text(fig.figcaption), "", "", "")

    process = main.select_one("#process")
    owner_map = {"us": "highlights", "city": "city", "comed": "comed"}
    page.update(process_heading=text(process.h2), process_lede=text(process.select_one("p.lede")))
    for li in process.select(".nx-steps li"):
        owners = " || ".join(owner_map[next(c for c in span["class"] if c != "nx-who")] for span in li.select(".nx-who"))
        add("process_step", text(li.h3), text(li.p), owners, "", "", "")

    page.update(areas_callout=text(main.select_one("#areas .nx-callout")))

    other = main.select_one("#other-services")
    page.update(feature_tag=text(other.select_one(".svc-feature-tag")), feature_title=text(other.select_one(".svc-feature h3")), feature_desc=text(other.select_one(".svc-feature p")))
    for link in other.select(".svc-mini"):
        add("other_service", text(link.b), text(link.span), link["href"], "", "", "")

    pricing = main.select_one("#pricing")
    tables = pricing.select("table")
    drivers = pricing.select_one(".nx-block")
    page.update(pricing_heading=strip_area(text(pricing.h2)), pricing_lede=text(pricing.select_one("p.lede")), pricing_caption=text(tables[0].caption),
                pricing_columns=" || ".join(text(th) for th in tables[0].select("thead th")), pricing_note=text(pricing.select_one("p.section-note")),
                drivers_heading=text(drivers.h3), incentives_note=text(pricing.select_one(".nx-callout")))
    for row in tables[0].select("tbody tr"):
        cells = [text(td) for td in row.find_all("td")]
        add("price_row", *(cells + [""] * 6)[:6])
    for row in drivers.select("tbody tr"):
        cells = [text(td) for td in row.find_all("td")]
        add("cost_driver", *(cells + [""] * 6)[:6])

    incl = main.select_one("#whats-included")
    inn, out = incl.select_one(".nx-in"), incl.select_one(".nx-out")
    hand, warranty = incl.select_one(".nx-ticks").find_parent(class_="nx-block"), incl.select_one(".nx-warranty").find_parent(class_="nx-block")
    page.update(incl_heading=text(incl.h2), incl_lede=text(incl.select_one("p.lede")),
                in_scope_heading=text(inn.h3), in_scope=" || ".join(text(li) for li in inn.select("li")), in_scope_note=text(inn.select_one("p.small")),
                separate_heading=text(out.h3), separate=" || ".join(text(li) for li in out.select("li")), separate_note=text(out.select_one("p.small")),
                handover_heading=text(hand.h3), handover=" || ".join(text(li) for li in hand.select("li")), warranty_heading=text(warranty.h3))
    for dt in warranty.select("dt"):
        add("warranty", text(dt), text(dt.find_next_sibling("dd")), "", "", "", "")

    local_faqs = []
    for details in main.select("#faq details"):
        question, answer = text(details.summary), text(details.select_one(".faq-body") or details.p)
        if question == "Do you cover my part of Chicago?":
            local_faqs.append((question, answer))
        else:
            add("faq", question, answer, "", "", "", "")

    cta = main.select_one(".closing-cta-section")
    page.update(cta_heading=strip_area(text(cta.h2)), cta_body=text(cta.select_one(".cta-final > p")),
                ready_heading=text(cta.select_one(".nx-ready h3")), ready_items=" || ".join(text(li) for li in cta.select(".nx-ready li")))

    shared = []
    for index, panel in enumerate(main.select("#guides .guide-panel")):
        title = text(panel.select_one(".guide-mobile-head"))
        body = panel.select_one(".guide-body")
        heading = text(body.h3)
        paragraphs = [p for p in body.find_all("p")]
        if index < 3:
            for p in paragraphs:
                shared.append([title, heading, "", text(p)])
            continue
        page.update(guide_title=title, guide_heading=heading, guide_intro=text(paragraphs[0]))
        sections = []
        for p in paragraphs[1:]:
            if p.b and text(p.b) == text(p):
                sections.append((text(p), []))
            elif sections:
                sections[-1][1].append(text(p))
        for heading_text, body_paragraphs in sections:
            add("guide_section", heading_text, "\n\n".join(body_paragraphs), "", "", "", "")

    trust = {
        "metrics": [(text(cell.b), text(cell.span)) for cell in main.select("#trust .trust-cell") if "google-proof-cell" not in cell.get("class", [])],
        "cards": [(text(card.h3), text(card.p)) for card in main.select("#trust .trust-cards article")],
        "lines": [text(item).lstrip("◆ ").strip() for item in hero.select(".trust-item")],
        "credits": {text(chip.b): text(chip.select_one(".area-credit")).replace("Photo: ", "") for chip in main.select("#areas .area-chip")},
        "photos": {text(chip.b): chip.img["src"] if chip.img else "" for chip in main.select("#areas .area-chip")},
        "disclaimer": text(main.select_one("#reviews .review-note")),
    }
    return page, items, local_faqs, shared, trust


def build(template, source_path, out):
    source = json.load(open(source_path, encoding="utf-8"))
    page, items, local_faqs, shared, trust = parse_template(template, source)
    area = source["area"][0]
    first = source["page"][0]

    wb = Workbook()
    readme = wb.active
    readme.title = "00 README"
    readme_rows = [
        ["Highlights Chicago: Service Pages V2 data workbook"],
        [""],
        ["What this is", "The content source for service pages on the new V2 template. Each sheet maps to a Sanity field (see 05 Field Map)."],
        ["01 Pages", "One row per service page. Single values and simple lists (items separated by ||)."],
        ["02 Items", "Every repeating card, table row, review, photo and FAQ. One row per item: owner (service_id, area:<slug> or settings), block, order, then columns a to f. The Field Map names what a to f mean for each block."],
        ["03 Area", "Shared area content (Chicago). Its neighbourhoods, shared library guides and area FAQs are in 02 Items with owner area:chicago."],
        ["04 Settings", "Company facts used on every page. Trust lines, metrics and cards are in 02 Items with owner settings."],
        ["05 Field Map", "Column or block -> Sanity field -> page section, with the writing rule for each."],
        [""],
        ["Example row", "Service 302 (Solar) is a complete example copied from the approved design file. Keep it as a reference or delete its rows before importing."],
        ["Placeholders", "Write [CONFIRM: question] where a fact needs Highlights to confirm it. The page shows it as a yellow chip and the build lists every one. Clear them before publishing."],
        ["Publish path", "pnpm content:extract -> pnpm content:build -> check http://localhost:3008/services/<slug> -> pnpm content:push (drafts) -> review in Studio -> pnpm content:push --publish"],
        ["Link anchors", ", ".join(ANCHORS)],
    ]
    for row in readme_rows:
        readme.append(row)
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 18
    readme.column_dimensions["B"].width = 140
    for row in readme.iter_rows(min_row=3):
        row[0].font = Font(bold=True)
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    def header(ws, cols, required=()):
        ws.append(cols)
        for index, cell in enumerate(ws[1], start=1):
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = REQ_FILL if cols[index - 1] in required else HEADER_FILL
            ws.column_dimensions[get_column_letter(index)].width = 28
        ws.freeze_panes = "B2"

    pages = wb.create_sheet("01 Pages")
    cols = [c[0] for c in PAGE_COLUMNS]
    header(pages, cols, {c[0] for c in PAGE_COLUMNS if c[3]})
    missing = [c for c in cols if c not in page]
    if missing:
        raise SystemExit(f"Example parse missed columns: {missing}")
    pages.append([page[c] for c in cols])

    itm = wb.create_sheet("02 Items")
    header(itm, ITEM_COLUMNS)
    for col, width in zip("ABCDEFGHI", [12, 16, 7, 40, 50, 40, 30, 22, 22]):
        itm.column_dimensions[col].width = width
    for item in items:
        itm.append(item[:9])
    for order, (name, note, url) in enumerate([s.split("::") for s in area["sub_areas"].split("||")], start=1):
        name = name.strip()
        itm.append(["area:chicago", "sub_area", order, name, note.strip(), trust["photos"].get(name) or url.strip(), trust["credits"].get(name, ""), "", ""])
    for order, row in enumerate(shared, start=1):
        itm.append(["area:chicago", "shared_guide", order, *row, "", ""])
    for order, (question, answer) in enumerate(local_faqs, start=1):
        itm.append(["area:chicago", "local_faq", order, question, answer, "", "", "", ""])
    for order, line in enumerate(trust["lines"], start=1):
        itm.append(["settings", "trust_line", order, line, "", "", "", "", ""])
    for order, (value, label) in enumerate(trust["metrics"], start=1):
        itm.append(["settings", "trust_metric", order, value, label, "", "", "", ""])
    for order, (title, body) in enumerate(trust["cards"], start=1):
        itm.append(["settings", "trust_card", order, title, body, "", "", "", ""])

    ar = wb.create_sheet("03 Area")
    header(ar, [c[0] for c in AREA_COLUMNS])
    ar.append([area.get(c[0], "").replace("||", " || ") if c[0] == "building_types" else area.get(c[0], "") for c in AREA_COLUMNS])

    st = wb.create_sheet("04 Settings")
    header(st, ["field", "value"])
    st.column_dimensions["A"].width, st.column_dimensions["B"].width = 26, 110
    values = {
        "company_name": first["company_name"], "site_url": first["site_url"], "phone_display": first["phone_display"], "phone_e164": first["phone_e164"],
        "email": first["email"], "address_street": first["address_street"], "address_city": first["address_city"], "address_state": first["address_state"],
        "address_zip": first["address_zip"], "shop_lat": first["shop_lat"], "shop_lng": first["shop_lng"], "schema_business_type": first["schema_business_type"],
        "google_rating": first["google_rating"], "google_review_count": first["google_review_count"], "google_reviews_url": first["google_reviews_url"],
        "google_verified_at": "2026-08-27", "trust_heading": first["trust_heading"], "trust_lede": first["trust_lede"],
        "reviews_heading": first["reviews_heading"], "reviews_disclaimer": trust["disclaimer"] or first.get("reviews_disclaimer", ""),
        "form_subtitle": first["form_subtitle"], "form_note": first["form_note"],
    }
    for field, _ in SETTINGS_FIELDS:
        st.append([field, values.get(field, "")])

    fm = wb.create_sheet("05 Field Map")
    header(fm, ["sheet", "column or block", "item columns (a to f)", "Sanity field", "page section", "required", "writing rule"])
    for col, width in zip("ABCDEFG", [12, 26, 48, 38, 22, 9, 80]):
        fm.column_dimensions[col].width = width
    for col, field, section, req, rule in PAGE_COLUMNS:
        fm.append(["01 Pages", col, "", f"v2ServicePage.{field}", section, "yes" if req else "", rule])
    for owner, blocks, doc in (("<service_id>", PAGE_BLOCKS, "v2ServicePage"), ("area:<slug>", AREA_BLOCKS, "v2ServiceArea"), ("settings", SETTINGS_BLOCKS, "v2SiteSettings")):
        for block, (fields, target, section, rule) in blocks.items():
            mapping = ", ".join(f"{letter}={name}" for letter, name in zip("abcdef", fields))
            fm.append(["02 Items", f"{block}  (owner {owner})", mapping, f"{doc}.{target}", section, "", rule])
    for col, field in AREA_COLUMNS:
        fm.append(["03 Area", col, "", f"v2ServiceArea.{field}", "Every page", "yes", ""])
    for col, field in SETTINGS_FIELDS:
        fm.append(["04 Settings", col, "", f"v2SiteSettings.{field}", "Every page", "yes", ""])

    for ws in (pages, itm, ar, st, fm):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.auto_filter.ref = ws.dimensions
    for cell in pages[2]:
        cell.fill = EXAMPLE_FILL

    wb.save(out)
    print(f"Wrote {out}: 1 example page, {len(items)} page items, {len(shared)} shared guide rows")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--source", required=True)
    parser.add_argument("--out", default=DEFAULT_OUT)
    args = parser.parse_args()
    build(args.template, args.source, args.out)
