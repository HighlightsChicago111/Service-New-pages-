"""Read the V2 workbook into data/v2-source.json (no network, no Sanity).

Usage: python -I scripts/extract_workbook.py ["<workbook>.xlsx"]
"""

import json
import os
import sys
from datetime import date, datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook  # noqa: E402

from workbook_spec import AREA_COLUMNS, ITEM_COLUMNS, PAGE_COLUMNS, SETTINGS_FIELDS  # noqa: E402

DEFAULT = "Highlights Chicago - Service Pages V2 Data.xlsx"
OUT = os.path.join("data", "v2-source.json")


def clean(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def rows(ws, expected):
    header = [clean(c.value) for c in ws[1]]
    missing = [c for c in expected if c not in header]
    extra = [c for c in header if c and c not in expected]
    if missing or extra:
        raise SystemExit(f"Sheet '{ws.title}': missing columns {missing}, unknown columns {extra}. Compare with 05 Field Map.")
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = {header[i]: clean(v) for i, v in enumerate(row) if i < len(header) and header[i]}
        if any(record.values()):
            out.append(record)
    return out


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    wb = load_workbook(path, data_only=True)
    for name in ("01 Pages", "02 Items", "03 Area", "04 Settings"):
        if name not in wb.sheetnames:
            raise SystemExit(f"Workbook has no '{name}' sheet")
    settings_rows = rows(wb["04 Settings"], ["field", "value"])
    settings = {r["field"]: r["value"] for r in settings_rows}
    unknown = [k for k in settings if k not in {f for f, _ in SETTINGS_FIELDS}]
    if unknown:
        raise SystemExit(f"04 Settings has unknown fields: {unknown}")
    data = {
        "workbook": os.path.basename(path),
        "pages": rows(wb["01 Pages"], [c[0] for c in PAGE_COLUMNS]),
        "items": rows(wb["02 Items"], ITEM_COLUMNS),
        "areas": rows(wb["03 Area"], [c[0] for c in AREA_COLUMNS]),
        "settings": settings,
    }
    os.makedirs("data", exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as handle:
        json.dump(data, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(f"Wrote {OUT}: {len(data['pages'])} pages, {len(data['items'])} items, {len(data['areas'])} areas")


if __name__ == "__main__":
    main()
